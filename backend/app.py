import re
from pathlib import Path
from flask import Flask, jsonify, request, send_from_directory
from docx import Document

ROOT_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIR = ROOT_DIR / 'frontend'
IMAGES_DIR = ROOT_DIR / 'images'
FEEDBACK_DIR = ROOT_DIR / 'Feedback'
FEEDBACK_DIR.mkdir(parents=True, exist_ok=True)

app = Flask(__name__, static_folder=str(FRONTEND_DIR), static_url_path='')

SCHOOL = {
    "name": "T.B.Jaya Zahira Vidyalaya",
    "founding_year": 1921,
    "motto": "THE LEARNED IS ONE WHO PRACTICES WHAT HE KNOWS",
    "medium": ["SINHALA", "TAMIL"],
    "students_approx": 700,
    "teachers_approx": 60,
    "principal": "MR. M.R.M.MANZIL",
    "website_creator": "M.I.MUHAMMADH (2025 O/L BATCH)",
    "address": "24, SIR JAMES PEIRIS MAWATHA, COLOMBO 02",
    "phone": "0112 300 338",
    "email": "tbjzc2013@gmail.com",
    "facebook": "https://web.facebook.com/T.BJZAHIRAOFFICIAL/?_rdc=1&_rdr#",
    "colors": ["MEROON", "YELLOW", "GREEN"],
    "houses": [
        {"name": "RAFAH", "color": "GREEN"},
        {"name": "CAIRO", "color": "YELLOW"},
        {"name": "JUMAIL", "color": "MEROON"}
    ],
    "sports_clubs": ["FOOTBALL", "CRICKET", "HOCKEY", "KARATE", "CHESS"],
    "achievements": ["All island first places in Islamic cultural competitions", "2024 O/L 9A student"]
}

CLUBS = {
    "Sinhala Literary Association": {
        "teacher_in_charge": "MRS. SRIYANI EKANAYAKE",
        "president": "M.A. ABDUL AZEEZ",
        "secretary": "M.I.MUHAMMADH",
        "vice_president": "MAHDI RIZWI",
    },
    "Tamil Literary Association": {
        "teacher_in_charge": "MRS. M. M. F. MINHA",
        "president": "M.I.A.ASLAM",
        "secretary": "M.AKAM",
    },
    "English Literary Association": {
        "teacher_in_charge": "MRS. FAZREEN MOHAMED",
        "president": "M.RM.ALFATH",
    },
    # Additional clubs truncated for brevity; full data can be added as needed
}


@app.route('/')
def frontpage():
    return app.send_static_file('index.html')


@app.route('/images/<path:filename>')
def serve_image(filename):
    return send_from_directory(str(IMAGES_DIR), filename)


@app.route('/info')
def info():
    return jsonify(SCHOOL)


@app.route('/clubs')
def clubs():
    q = request.args.get('q', '').lower()
    if q:
        filtered = {k: v for k, v in CLUBS.items() if q in k.lower()}
        return jsonify(filtered)
    return jsonify(CLUBS)


@app.route('/submit-feedback', methods=['POST'])
def submit_feedback():
    try:
        data = request.get_json(silent=True) or {}
        name = (data.get('name') or '').strip()
        email = (data.get('email') or '').strip()
        message = (data.get('message') or '').strip()

        if not name or not email or not message:
            return jsonify(error='Name, email, and feedback message are required.'), 400

        if not re.match(r'^[^@\s]+@[^@\s]+\.[^@\s]+$', email):
            return jsonify(error='A valid email address is required.'), 400

        if len(message.split()) > 700:
            return jsonify(error='Feedback message must be 700 words or fewer.'), 400

        feedback_path = FEEDBACK_DIR / 'Feedback.docx'
        if feedback_path.exists():
            document = Document(str(feedback_path))
            document.add_paragraph('')
        else:
            document = Document()

        document.add_paragraph('Feedback')
        document.add_paragraph(f'Name - {name}')
        document.add_paragraph(f'Email - {email}')
        document.add_paragraph('______________________________')
        document.add_paragraph('Full details:')
        for paragraph in message.splitlines():
            document.add_paragraph(paragraph)
        document.save(str(feedback_path))

        return jsonify(message='Feedback saved successfully.')
    except Exception as exc:
        return jsonify(error=str(exc) or 'An internal error occurred while saving feedback.'), 500


@app.route('/search')
def search():
    q = request.args.get('q', '').lower()
    results = {k: v for k, v in CLUBS.items() if q in k.lower()} if q else {}
    return jsonify(results)


if __name__ == '__main__':
    app.run(host='0.0.0.0', port=5000, debug=True)
